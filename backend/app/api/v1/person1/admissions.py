import csv
import io
import re
from datetime import datetime
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.dependencies import get_db
from app.models.academic_year import AcademicYear
from app.models.admission import Admission
from app.models.class_section import Class, Section
from app.models.parent import Parent
from app.models.student import Student
from app.schemas.admission import (
    AdmissionRegistrationCreate,
    AdmissionRegistrationRead,
    BulkAdmissionError,
    BulkAdmissionImportResponse,
    BulkAdmissionImportSummary,
    BulkAdmissionPreviewRow,
    BulkAdmissionValidationResponse,
)

router = APIRouter(prefix="/admissions", tags=["person1-admissions"])

MAX_BULK_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_BULK_ROWS = 200

BULK_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("admission_date", "Admission Date", True, "Use YYYY-MM-DD, for example 2026-10-11."),
    ("academic_year", "Academic Year", True, "Use the school year label, for example 2026-27."),
    ("class_admitted", "Class Admitted", True, "Use values such as Nursery, Prep, Class 1, Class 6."),
    ("section", "Section", True, "Use values such as A, B, C, D, or E."),
    ("admission_type", "Admission Type", True, "Use New Admission or Transfer Admission."),
    ("student_name", "Student Name", True, "Enter the student's full name."),
    ("date_of_birth", "Date of Birth", True, "Use YYYY-MM-DD, for example 2018-04-12."),
    ("gender", "Gender", True, "Use Male, Female, Other, or Prefer not to say."),
    ("registration_number", "Registration Number", False, "Optional. Format as text to preserve leading zeros."),
    ("father_name", "Father Name", False, "Optional."),
    ("mother_name", "Mother Name", False, "Optional."),
    ("guardian_name", "Guardian Name", False, "Optional."),
    ("parent_phone", "Parent Phone", False, "Optional 10-digit phone number."),
    ("parent_email", "Parent Email", False, "Optional email address."),
    ("parent_address", "Parent Address", False, "Optional address."),
    ("previous_school", "Previous School", False, "Required for Transfer Admission."),
    ("previous_class", "Previous Class", False, "Required for Transfer Admission."),
    ("previous_admission_number", "Previous Admission Number", False, "Optional."),
    ("previous_tc_number", "Previous TC Number", False, "Required for Transfer Admission."),
]

REQUIRED_BULK_COLUMNS = {column for column, _, required, _ in BULK_COLUMNS if required}
OPTIONAL_BULK_COLUMNS = {column for column, _, required, _ in BULK_COLUMNS if not required}
ALLOWED_BULK_COLUMNS = REQUIRED_BULK_COLUMNS | OPTIONAL_BULK_COLUMNS
ADMISSION_TYPES = {"New Admission", "Transfer Admission"}
GENDERS = {"Male", "Female", "Other", "Prefer not to say"}


@router.get("")
def list_admissions() -> dict[str, str]:
    return {"message": "Admissions endpoint placeholder"}


@router.post("", status_code=status.HTTP_201_CREATED)
def create_admission(payload: AdmissionRegistrationCreate, db: Session = Depends(get_db)) -> AdmissionRegistrationRead:
    now = datetime.utcnow()
    try:
        registration = _create_admission_registration(db, payload, now)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Admission could not be saved because a generated identifier already exists.") from exc

    return registration


@router.get("/bulk/template")
def download_bulk_admission_template(format: str = Query(default="xlsx", pattern="^(xlsx|csv)$")) -> StreamingResponse:
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([label for _, label, _, _ in BULK_COLUMNS])
        writer.writerow(["2026-10-11", "2026-27", "Class 1", "A", "New Admission", "Sample Student", "2018-04-12", "Female", "REG-2026-001"])
        writer.writerow([])
        writer.writerow(["Column", "Required", "Instructions"])
        for _, label, required, instructions in BULK_COLUMNS:
            writer.writerow([label, "Yes" if required else "No", instructions])
        content = output.getvalue().encode("utf-8-sig")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="text/csv",
            headers={"Content-Disposition": 'attachment; filename="bulk-admission-template.csv"'},
        )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="Excel template support requires openpyxl to be installed.") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Students"
    worksheet.append([label for _, label, _, _ in BULK_COLUMNS])
    worksheet.append(["2026-10-11", "2026-27", "Class 1", "A", "New Admission", "Sample Student", "2018-04-12", "Female", "REG-2026-001"])
    header_fill = PatternFill("solid", fgColor="DBEAFE")
    required_fill = PatternFill("solid", fgColor="FEE2E2")
    for index, (_, _, required, _) in enumerate(BULK_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = required_fill if required else header_fill
        worksheet.column_dimensions[cell.column_letter].width = 24

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Column", "Required", "Instructions"])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for _, label, required, text in BULK_COLUMNS:
        instructions.append([label, "Yes" if required else "No", text])
    instructions.append([])
    instructions.append(["Note", "Documents and photos cannot be imported from local spreadsheet paths. Use the existing document upload workflow after admission."])
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 12
    instructions.column_dimensions["C"].width = 90

    output_bytes = io.BytesIO()
    workbook.save(output_bytes)
    output_bytes.seek(0)
    return StreamingResponse(
        output_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="bulk-admission-template.xlsx"'},
    )


@router.post("/bulk/validate")
async def validate_bulk_admissions(file: UploadFile = File(...), db: Session = Depends(get_db)) -> BulkAdmissionValidationResponse:
    content = await _read_bulk_upload(file)
    validation, _ = _validate_bulk_content(db, file.filename or "", content)
    return validation


@router.post("/bulk/import")
async def import_bulk_admissions(file: UploadFile = File(...), db: Session = Depends(get_db)) -> BulkAdmissionImportResponse:
    content = await _read_bulk_upload(file)
    validation, payloads_by_row = _validate_bulk_content(db, file.filename or "", content)
    valid_payloads = [
        payloads_by_row[row.row_number]
        for row in validation.rows
        if row.status == "valid" and row.row_number in payloads_by_row
    ]

    imported_records: list[AdmissionRegistrationRead] = []
    failed_rows = 0
    now = datetime.utcnow()
    admission_counts: dict[int, int] = {}
    student_counts: dict[int, int] = {}

    if valid_payloads:
        try:
            for payload in valid_payloads:
                imported_records.append(_create_admission_registration(db, payload, now, admission_counts, student_counts))
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            failed_rows = len(valid_payloads)
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Bulk admission import failed because one or more records conflicted with existing data.") from exc
        except Exception as exc:
            db.rollback()
            failed_rows = len(valid_payloads)
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Bulk admission import failed. No student records were saved.") from exc

    summary = BulkAdmissionImportSummary(
        rows_processed=validation.total_rows,
        imported=len(imported_records),
        invalid_rows=validation.invalid_rows,
        duplicate_rows=validation.duplicate_rows,
        failed_rows=failed_rows,
        academic_years_created=validation.academic_years_to_create if imported_records else [],
        classes_created=validation.classes_to_create if imported_records else [],
        sections_created=validation.sections_to_create if imported_records else [],
    )
    return BulkAdmissionImportResponse(summary=summary, imported_records=imported_records, validation=validation)


def _create_admission_registration(
    db: Session,
    payload: AdmissionRegistrationCreate,
    now: datetime,
    admission_counts: dict[int, int] | None = None,
    student_counts: dict[int, int] | None = None,
) -> AdmissionRegistrationRead:
    academic_year = _get_or_create_academic_year(db, payload.academic_year, now)
    class_record = _get_or_create_class(db, payload.class_admitted)
    section_record = _get_or_create_section(db, payload.section)
    admission_number = _next_admission_number(db, payload.admission_date.year, admission_counts)
    student_code = _next_student_code(db, payload.admission_date.year, student_counts)

    student = Student(
        student_id=student_code,
        admission_number=admission_number,
        name=payload.student_name.strip(),
        date_of_birth=payload.date_of_birth,
        gender=payload.gender,
        class_id=class_record.id,
        section_id=section_record.id,
        status="active",
        created_at=now,
        updated_at=now,
    )
    db.add(student)
    db.flush()

    admission = Admission(
        student_id=student.id,
        admission_number=admission_number,
        registration_number=payload.registration_number,
        admission_date=payload.admission_date,
        academic_year_id=academic_year.id,
        class_id=class_record.id,
        section_id=section_record.id,
        admission_type=payload.admission_type,
        previous_school=payload.previous_school,
        previous_class=payload.previous_class,
        previous_admission_number=payload.previous_admission_number,
        previous_tc_number=payload.previous_tc_number,
        created_at=now,
        updated_at=now,
    )
    db.add(admission)
    db.flush()

    if any([payload.father_name, payload.mother_name, payload.guardian_name, payload.parent_phone, payload.parent_email, payload.parent_address]):
        db.add(
            Parent(
                student_id=student.id,
                father_name=payload.father_name,
                mother_name=payload.mother_name,
                guardian_name=payload.guardian_name,
                phone=payload.parent_phone,
                email=payload.parent_email,
                address=payload.parent_address,
                created_at=now,
                updated_at=now,
            )
        )
        db.flush()

    return AdmissionRegistrationRead(
        student_id=student.id,
        student_code=student.student_id,
        admission_id=admission.id,
        admission_number=admission.admission_number or admission_number,
        student_name=student.name,
        academic_year=academic_year.name,
        class_name=class_record.name,
        section_name=section_record.name,
    )


def _get_or_create_academic_year(db: Session, name: str, now: datetime) -> AcademicYear:
    academic_year = db.execute(select(AcademicYear).where(func.lower(AcademicYear.name) == name.strip().lower())).scalar_one_or_none()
    if academic_year is None:
        academic_year = AcademicYear(name=name.strip(), created_at=now, updated_at=now)
        db.add(academic_year)
        db.flush()
    return academic_year


def _get_or_create_class(db: Session, name: str) -> Class:
    class_record = db.execute(select(Class).where(func.lower(Class.name) == name.strip().lower())).scalar_one_or_none()
    if class_record is None:
        class_record = Class(name=name.strip())
        db.add(class_record)
        db.flush()
    return class_record


def _get_or_create_section(db: Session, name: str) -> Section:
    section = db.execute(select(Section).where(func.lower(Section.name) == name.strip().lower())).scalar_one_or_none()
    if section is None:
        section = Section(name=name.strip())
        db.add(section)
        db.flush()
    return section


async def _read_bulk_upload(file: UploadFile) -> bytes:
    filename = file.filename or ""
    if not filename.lower().endswith((".csv", ".xlsx")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only .xlsx and .csv files are supported.")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file is empty.")
    if len(content) > MAX_BULK_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Bulk admission files must be 2 MB or smaller.")
    return content


def _validate_bulk_content(db: Session, filename: str, content: bytes) -> tuple[BulkAdmissionValidationResponse, dict[int, AdmissionRegistrationCreate]]:
    raw_rows = _parse_bulk_rows(filename, content)
    errors: list[BulkAdmissionError] = []
    payloads_by_row: dict[int, AdmissionRegistrationCreate] = {}
    preview_rows: list[BulkAdmissionPreviewRow] = []

    if not raw_rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file does not contain any student rows.")
    if len(raw_rows) > MAX_BULK_ROWS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Bulk admission files can contain at most {MAX_BULK_ROWS} student rows.")

    header_columns = set(raw_rows[0][1].keys()) if raw_rows else set()
    missing_columns = sorted(REQUIRED_BULK_COLUMNS - header_columns)
    for column in missing_columns:
        errors.append(
            BulkAdmissionError(
                row_number=1,
                column=_column_label(column),
                message=f"Required column '{_column_label(column)}' is missing.",
                suggestion="Download the latest template and keep all required headers.",
            )
        )

    existing_academic_years = _existing_names(db, AcademicYear)
    existing_classes = _existing_names(db, Class)
    existing_sections = _existing_names(db, Section)
    academic_years_to_create: set[str] = set()
    classes_to_create: set[str] = set()
    sections_to_create: set[str] = set()
    seen_registration_numbers: dict[str, int] = {}
    seen_student_keys: dict[str, int] = {}

    for row_number, raw_row in raw_rows:
        row_errors: list[BulkAdmissionError] = []
        normalized = {column: _clean_cell(raw_row.get(column)) for column in ALLOWED_BULK_COLUMNS}

        for column in REQUIRED_BULK_COLUMNS:
            if not normalized.get(column):
                row_errors.append(
                    BulkAdmissionError(
                        row_number=row_number,
                        column=_column_label(column),
                        value=None,
                        message=f"{_column_label(column)} is required.",
                        suggestion="Fill this cell before uploading again.",
                    )
                )

        normalized["admission_date"] = _parse_date_cell(normalized.get("admission_date"), row_number, "admission_date", row_errors)
        normalized["date_of_birth"] = _parse_date_cell(normalized.get("date_of_birth"), row_number, "date_of_birth", row_errors)

        if isinstance(normalized["date_of_birth"], date) and normalized["date_of_birth"] > date.today():
            row_errors.append(
                BulkAdmissionError(
                    row_number=row_number,
                    column="Date of Birth",
                    value=normalized["date_of_birth"].isoformat(),
                    message="Date of birth cannot be in the future.",
                    suggestion="Enter the student's actual date of birth.",
                )
            )

        if normalized.get("admission_type") and normalized["admission_type"] not in ADMISSION_TYPES:
            row_errors.append(
                BulkAdmissionError(
                    row_number=row_number,
                    column="Admission Type",
                    value=str(normalized["admission_type"]),
                    message="Admission type is not supported.",
                    suggestion="Use New Admission or Transfer Admission.",
                )
            )

        if normalized.get("gender") and normalized["gender"] not in GENDERS:
            row_errors.append(
                BulkAdmissionError(
                    row_number=row_number,
                    column="Gender",
                    value=str(normalized["gender"]),
                    message="Gender is not supported.",
                    suggestion="Use Male, Female, Other, or Prefer not to say.",
                )
            )

        if normalized.get("parent_phone") and not re.fullmatch(r"\d{10}", str(normalized["parent_phone"])):
            row_errors.append(
                BulkAdmissionError(
                    row_number=row_number,
                    column="Parent Phone",
                    value=str(normalized["parent_phone"]),
                    message="Parent phone must be a 10-digit number.",
                    suggestion="Use digits only, for example 9000000099.",
                )
            )

        if normalized.get("parent_email") and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", str(normalized["parent_email"])):
            row_errors.append(
                BulkAdmissionError(
                    row_number=row_number,
                    column="Parent Email",
                    value=str(normalized["parent_email"]),
                    message="Parent email address is not valid.",
                    suggestion="Use an address such as parent@example.com or leave the cell blank.",
                )
            )

        if normalized.get("admission_type") == "Transfer Admission":
            for column in ("previous_school", "previous_class", "previous_tc_number"):
                if not normalized.get(column):
                    row_errors.append(
                        BulkAdmissionError(
                            row_number=row_number,
                            column=_column_label(column),
                            value=None,
                            message=f"{_column_label(column)} is required for transfer admissions.",
                            suggestion="Fill this value or change Admission Type to New Admission.",
                        )
                    )

        registration_number = str(normalized.get("registration_number") or "").strip()
        if registration_number:
            registration_key = registration_number.lower()
            previous_row = seen_registration_numbers.get(registration_key)
            if previous_row:
                row_errors.append(
                    BulkAdmissionError(
                        row_number=row_number,
                        column="Registration Number",
                        value=registration_number,
                        message=f"Registration number is duplicated in the uploaded file. It first appears on row {previous_row}.",
                        suggestion="Use a unique registration number or leave it blank.",
                    )
                )
            else:
                seen_registration_numbers[registration_key] = row_number

            if _registration_number_exists(db, registration_number):
                row_errors.append(
                    BulkAdmissionError(
                        row_number=row_number,
                        column="Registration Number",
                        value=registration_number,
                        message="Registration number already exists in saved admissions.",
                        suggestion="Use a unique registration number or verify this student was not already admitted.",
                    )
                )

        student_key = _student_duplicate_key(normalized)
        if student_key:
            previous_row = seen_student_keys.get(student_key)
            if previous_row:
                row_errors.append(
                    BulkAdmissionError(
                        row_number=row_number,
                        column="Student Name",
                        value=str(normalized.get("student_name") or ""),
                        message=f"This student appears to be duplicated in the uploaded file. The same name, date of birth, and parent phone first appear on row {previous_row}.",
                        suggestion="Remove the duplicate row or correct the identifying details.",
                    )
                )
            else:
                seen_student_keys[student_key] = row_number

            if _student_identity_exists(db, normalized):
                row_errors.append(
                    BulkAdmissionError(
                        row_number=row_number,
                        column="Student Name",
                        value=str(normalized.get("student_name") or ""),
                        message="A student with the same name, date of birth, and parent phone already exists.",
                        suggestion="Check the student register before importing.",
                    )
                )

        if normalized.get("academic_year") and str(normalized["academic_year"]).strip().lower() not in existing_academic_years:
            academic_years_to_create.add(str(normalized["academic_year"]).strip())
        if normalized.get("class_admitted") and str(normalized["class_admitted"]).strip().lower() not in existing_classes:
            classes_to_create.add(str(normalized["class_admitted"]).strip())
        if normalized.get("section") and str(normalized["section"]).strip().lower() not in existing_sections:
            sections_to_create.add(str(normalized["section"]).strip())

        payload = None
        if not row_errors and not errors:
            try:
                payload = AdmissionRegistrationCreate(**normalized)
                payloads_by_row[row_number] = payload
            except ValidationError as exc:
                row_errors.extend(_validation_errors(row_number, exc))

        row_status = "valid"
        if any("duplicate" in error.message.lower() or "already exists" in error.message.lower() for error in row_errors):
            row_status = "duplicate"
        elif row_errors or errors:
            row_status = "invalid"

        preview_rows.append(
            BulkAdmissionPreviewRow(
                row_number=row_number,
                status=row_status,
                admission_date=payload.admission_date if payload else normalized.get("admission_date") if isinstance(normalized.get("admission_date"), date) else None,
                academic_year=_string_or_none(normalized.get("academic_year")),
                class_admitted=_string_or_none(normalized.get("class_admitted")),
                section=_string_or_none(normalized.get("section")),
                admission_type=_string_or_none(normalized.get("admission_type")),
                registration_number=_string_or_none(normalized.get("registration_number")),
                student_name=_string_or_none(normalized.get("student_name")),
                date_of_birth=payload.date_of_birth if payload else normalized.get("date_of_birth") if isinstance(normalized.get("date_of_birth"), date) else None,
                gender=_string_or_none(normalized.get("gender")),
                parent_phone=_string_or_none(normalized.get("parent_phone")),
                errors=row_errors,
            )
        )

    invalid_rows = sum(1 for row in preview_rows if row.status == "invalid")
    duplicate_rows = sum(1 for row in preview_rows if row.status == "duplicate")
    valid_rows = sum(1 for row in preview_rows if row.status == "valid")
    validation = BulkAdmissionValidationResponse(
        total_rows=len(preview_rows),
        valid_rows=valid_rows,
        invalid_rows=invalid_rows,
        duplicate_rows=duplicate_rows,
        rows=preview_rows,
        errors=errors,
        academic_years_to_create=sorted(academic_years_to_create),
        classes_to_create=sorted(classes_to_create),
        sections_to_create=sorted(sections_to_create),
    )
    return validation, payloads_by_row


def _parse_bulk_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, object]]]:
    if filename.lower().endswith(".csv"):
        return _parse_csv_rows(content)
    if filename.lower().endswith(".xlsx"):
        return _parse_xlsx_rows(content)
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only .xlsx and .csv files are supported.")


def _parse_csv_rows(content: bytes) -> list[tuple[int, dict[str, object]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="CSV files must be saved as UTF-8.") from exc

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The CSV file is missing a header row.")

    header_map = {_normalize_header(header): header for header in reader.fieldnames if _normalize_header(header)}
    rows: list[tuple[int, dict[str, object]]] = []
    for index, row in enumerate(reader, start=2):
        if not any(_clean_cell(value) for value in row.values()):
            continue
        rows.append((index, {column: row.get(original_header) for column, original_header in header_map.items()}))
    return rows


def _parse_xlsx_rows(content: bytes) -> list[tuple[int, dict[str, object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="Excel upload support requires openpyxl to be installed.") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The Excel file could not be read. Please upload a valid .xlsx file.") from exc

    worksheet = workbook.active
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_values:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The Excel file is missing a header row.")

    header_map = {_normalize_header(value): index for index, value in enumerate(header_values) if _normalize_header(value)}
    rows: list[tuple[int, dict[str, object]]] = []
    for row_index, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(_clean_cell(value) for value in row_values):
            continue
        rows.append((row_index, {column: row_values[index] if index < len(row_values) else None for column, index in header_map.items()}))
    return rows


def _normalize_header(value: object) -> str:
    label = str(value or "").strip().lower()
    label = re.sub(r"\s*\(.*?\)\s*", " ", label)
    label = re.sub(r"[^a-z0-9]+", "_", label).strip("_")
    aliases = {
        "admission_date": "admission_date",
        "academic_year": "academic_year",
        "class": "class_admitted",
        "class_admitted": "class_admitted",
        "section": "section",
        "admission_type": "admission_type",
        "student_name": "student_name",
        "full_name": "student_name",
        "date_of_birth": "date_of_birth",
        "dob": "date_of_birth",
        "gender": "gender",
        "registration_number": "registration_number",
        "father_name": "father_name",
        "mother_name": "mother_name",
        "guardian_name": "guardian_name",
        "parent_phone": "parent_phone",
        "parent_guardian_phone": "parent_phone",
        "parent_email": "parent_email",
        "email": "parent_email",
        "parent_address": "parent_address",
        "previous_school": "previous_school",
        "previous_school_name": "previous_school",
        "previous_class": "previous_class",
        "previous_admission_number": "previous_admission_number",
        "previous_tc_number": "previous_tc_number",
        "previous_transfer_certificate_number": "previous_tc_number",
    }
    return aliases.get(label, label)


def _clean_cell(value: object) -> object | None:
    if value is None:
        return None
    if isinstance(value, str):
        trimmed = value.strip()
        return trimmed or None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return value


def _parse_date_cell(value: object, row_number: int, column: str, errors: list[BulkAdmissionError]) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    errors.append(
        BulkAdmissionError(
            row_number=row_number,
            column=_column_label(column),
            value=text,
            message=f"{_column_label(column)} must be a valid date.",
            suggestion="Use YYYY-MM-DD, DD-MM-YYYY, or DD/MM/YYYY.",
        )
    )
    return None


def _validation_errors(row_number: int, exc: ValidationError) -> list[BulkAdmissionError]:
    result: list[BulkAdmissionError] = []
    for error in exc.errors():
        column = str(error.get("loc", ["Unknown"])[0])
        result.append(
            BulkAdmissionError(
                row_number=row_number,
                column=_column_label(column),
                value=None,
                message=str(error.get("msg", "Value is invalid.")),
                suggestion="Check the value against the template instructions.",
            )
        )
    return result


def _column_label(column: str) -> str:
    for key, label, _, _ in BULK_COLUMNS:
        if key == column:
            return label
    return column.replace("_", " ").title()


def _string_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _existing_names(db: Session, model: type[AcademicYear] | type[Class] | type[Section]) -> set[str]:
    return {name.lower() for name in db.execute(select(model.name)).scalars().all() if name}


def _registration_number_exists(db: Session, registration_number: str) -> bool:
    return (
        db.execute(
            select(Admission.id).where(func.lower(func.coalesce(Admission.registration_number, "")) == registration_number.strip().lower())
        ).first()
        is not None
    )


def _student_duplicate_key(row: dict[str, object]) -> str | None:
    name = _string_or_none(row.get("student_name"))
    dob = row.get("date_of_birth")
    phone = _string_or_none(row.get("parent_phone"))
    if not name or not isinstance(dob, date) or not phone:
        return None
    return f"{name.strip().lower()}|{dob.isoformat()}|{phone.strip().lower()}"


def _student_identity_exists(db: Session, row: dict[str, object]) -> bool:
    name = _string_or_none(row.get("student_name"))
    dob = row.get("date_of_birth")
    phone = _string_or_none(row.get("parent_phone"))
    if not name or not isinstance(dob, date) or not phone:
        return False
    return (
        db.execute(
            select(Student.id)
            .join(Parent, Parent.student_id == Student.id)
            .where(
                func.lower(Student.name) == name.strip().lower(),
                Student.date_of_birth == dob,
                Parent.phone == phone.strip(),
            )
        ).first()
        is not None
    )


def _next_admission_number(db: Session, year: int, counts: dict[int, int] | None = None) -> str:
    prefix = f"ADM/{year}/"
    if counts is None:
        count = db.execute(select(func.count()).select_from(Student).where(Student.admission_number.like(f"{prefix}%"))).scalar_one()
        return f"{prefix}{count + 1:04d}"

    if year not in counts:
        counts[year] = db.execute(select(func.count()).select_from(Student).where(Student.admission_number.like(f"{prefix}%"))).scalar_one()
    counts[year] += 1
    return f"{prefix}{counts[year]:04d}"


def _next_student_code(db: Session, year: int, counts: dict[int, int] | None = None) -> str:
    prefix = f"STU-{year}-"
    if counts is None:
        count = db.execute(select(func.count()).select_from(Student).where(Student.student_id.like(f"{prefix}%"))).scalar_one()
        return f"{prefix}{count + 1:04d}"

    if year not in counts:
        counts[year] = db.execute(select(func.count()).select_from(Student).where(Student.student_id.like(f"{prefix}%"))).scalar_one()
    counts[year] += 1
    return f"{prefix}{counts[year]:04d}"
